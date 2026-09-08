#!/usr/bin/env python3
"""Convert CHL-C-ROB08 scanned PDF posters into an editable Word document.

Brand names RobotArt / PQ Art are normalized to PQArt.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
import pymupdf
from PIL import Image


SRC_PDF = Path("/home/ubuntu/.cursor/projects/workspace/uploads/CHL-C-ROB08___________7df3.pdf")
if not SRC_PDF.exists():
    SRC_PDF = Path("/home/ubuntu/.cursor/projects/workspace/uploads/CHL-C-ROB08___________82d0.pdf")
OUT_DOCX = Path("/workspace/CHL-C-ROB08_工业机器人离线编程.docx")
JPEG_DIR = Path("/tmp/pdf_pages")

BLUE = RGBColor(0x1F, 0x5F, 0xA8)
DARK = RGBColor(0x22, 0x22, 0x22)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_FILL = "1F5FA8"
ALT_FILL = "E8F1FB"
SECTION_FILL = "D6E6F5"


def set_run_font(run, name="Calibri", east_asia="微软雅黑", size=11, bold=False, color=DARK):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), name)


def shade_cell(cell, fill: str):
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="B4C9DE"):
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    tc_borders = tcpr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tcpr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tc_borders.append(el)


def cell_text(cell, text, *, bold=False, size=10, color=DARK, fill=None, align="left"):
    cell.text = ""
    p = cell.paragraphs[0]
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    if fill:
        shade_cell(cell, fill)
    set_cell_borders(cell)
    cell.vertical_alignment = 1  # center


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    size = {0: 22, 1: 16, 2: 13, 3: 12}[level]
    set_run_font(run, size=size, bold=True, color=BLUE)
    return p


def add_body(doc, text, *, size=11, bold=False, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_resource_table(doc, rows):
    """rows: list of (section, name, types) for one chapter."""
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    hdr = table.rows[0].cells
    cell_text(hdr[0], "节次", bold=True, size=10, color=WHITE, fill=HEADER_FILL, align="center")
    cell_text(hdr[1], "资源名称", bold=True, size=10, color=WHITE, fill=HEADER_FILL, align="center")
    cell_text(hdr[2], "资源类型", bold=True, size=10, color=WHITE, fill=HEADER_FILL, align="center")

    for i, (section, name, types) in enumerate(rows):
        cells = table.add_row().cells
        fill = ALT_FILL if i % 2 == 0 else "FFFFFF"
        cell_text(cells[0], section, size=10, fill=fill)
        cell_text(cells[1], name, size=10, fill=fill)
        cell_text(cells[2], types, size=10, fill=fill, align="center")

    for row in table.rows:
        row.cells[0].width = Cm(5.5)
        row.cells[1].width = Cm(10.5)
        row.cells[2].width = Cm(2.8)

    # Repeat header row on each page
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    return table


def jpeg_page(src: Path, dest: Path, quality=88) -> Path:
    im = Image.open(src).convert("RGB")
    im.save(dest, "JPEG", quality=quality, optimize=True)
    return dest


def configure_a3_landscape(section):
    section.page_width = Cm(42.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(0.5)
    section.right_margin = Cm(0.5)
    section.top_margin = Cm(0.5)
    section.bottom_margin = Cm(0.5)
    section.header_distance = Cm(0.3)
    section.footer_distance = Cm(0.3)


def configure_a4_portrait(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)


def add_fullpage_poster(doc, image_path: Path):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    run.add_picture(str(image_path), height=Cm(28.5))


def extract_pdf_jpegs(pdf_path: Path):
    JPEG_DIR.mkdir(parents=True, exist_ok=True)
    doc = pymupdf.open(pdf_path)
    jpgs = []
    for i, page in enumerate(doc, start=1):
        pix = page.get_pixmap(matrix=pymupdf.Matrix(1, 1), alpha=False)
        png_path = JPEG_DIR / f"page{i}_rgb.png"
        pix.save(png_path)
        jpgs.append(jpeg_page(png_path, JPEG_DIR / f"page{i}.jpg"))
    return jpgs


def build():
    jpg1, jpg2 = extract_pdf_jpegs(SRC_PDF)

    doc = Document()

    # Pages 1-2: 1:1 visual conversion of the scanned poster PDF
    configure_a3_landscape(doc.sections[0])
    add_fullpage_poster(doc, jpg1)
    poster2 = doc.add_section()
    configure_a3_landscape(poster2)
    add_fullpage_poster(doc, jpg2)

    # Remaining pages: editable resource list
    body = doc.add_section()
    configure_a4_portrait(body)

    # styles
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("CHL-C-ROB08")
    set_run_font(r, size=14, bold=True, color=BLUE)

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h.paragraph_format.space_after = Pt(12)
    r = h.add_run("工业机器人离线编程")
    set_run_font(r, size=26, bold=True, color=BLUE)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(16)
    r = sub.add_run("可编辑资源清单（软件名称统一为 PQArt；前两页为原海报版式）")
    set_run_font(r, size=11, color=RGBColor(0x55, 0x55, 0x55))

    add_heading(doc, "课程概述", 1)
    add_body(
        doc,
        "课程基于华航唯实自主研发的PQArt软件开发，通过典型工作任务，结合知识点PPT、"
        "实操微课视频等教学资源，讲解工业机器人离线编程相关知识和技能。",
    )
    add_body(doc, "课程配套教材：工业机器人离线编程（PQArt）", bold=True, size=11)

    add_heading(doc, "资源类型说明", 1)
    legend = doc.add_table(rows=1, cols=2)
    hdr = legend.rows[0].cells
    cell_text(hdr[0], "类型", bold=True, size=10, color=WHITE, fill=HEADER_FILL, align="center")
    cell_text(hdr[1], "说明", bold=True, size=10, color=WHITE, fill=HEADER_FILL, align="center")
    for i, (k, v) in enumerate(
        [
            ("PDF", "文档课件"),
            ("PPT", "演示文稿"),
            ("视频", "实操 / 微课视频"),
            ("动画", "动画演示"),
        ]
    ):
        cells = legend.add_row().cells
        fill = ALT_FILL if i % 2 == 0 else "FFFFFF"
        cell_text(cells[0], k, size=10, fill=fill, align="center", bold=True)
        cell_text(cells[1], v, size=10, fill=fill)
    doc.add_paragraph()

    add_heading(doc, "资源清单", 1)

    add_heading(doc, "第一章 工业机器人离线编程概述", 2)
    add_resource_table(
        doc,
        [
            ("第1节 工业机器人离线编程应用", "工业机器人离线编程应用", "PPT、视频"),
            ("第2节 工业机器人常用的离线编程软件", "工业机器人常用的离线编程软件", "PPT、视频"),
            ("第2节 工业机器人常用的离线编程软件", "工业机器人离线编程软件PQArt", "视频"),
        ],
    )

    add_heading(doc, "第二章 离线编程软件开发环境介绍", 2)
    add_resource_table(
        doc,
        [
            ("第1节 工业机器人离线编程部署", "工业机器人离线编程部署", "PPT、视频"),
            ("第2节 工业机器人离线编程仿真软件PQArt的界面介绍", "PQArt离线编程仿真软件界面", "PPT"),
            ("第2节 工业机器人离线编程仿真软件PQArt的界面介绍", "PQArt离线编程仿真软件界面", "视频"),
            ("第3节 工业机器人离线编程仿真软件三维球的基本操作", "三维球仿真软件的基本操作", "PPT、视频"),
            ("第3节 工业机器人离线编程仿真软件三维球的基本操作", "物块装配-PQArt三维球基本操作练习（1）", "视频"),
            ("第3节 工业机器人离线编程仿真软件三维球的基本操作", "物块装配-PQArt三维球基本操作练习（2）", "视频"),
            ("第3节 工业机器人离线编程仿真软件三维球的基本操作", "物块装配动画", "动画"),
            ("第4节 工业机器人离线编程仿真软件PQArt的基本操作", "工业机器人离线编程仿真软件的基本操作", "PPT、视频"),
            ("第4节 工业机器人离线编程仿真软件PQArt的基本操作", "工业机器人离线编程软件PQArt入门基本操作", "视频"),
        ],
    )

    add_heading(doc, "第三章 工业机器人工作站系统构建", 2)
    add_resource_table(
        doc,
        [
            ("第1节 导入机器人", "工业机器人工作站系统构建 - 导入机器人", "PPT、视频"),
            ("第2节 导入工件", "工业机器人工作站系统构建 - 导入工件", "PPT、视频"),
            ("第3节 导入工具", "工业机器人工作站系统构建 - 导入工具", "PPT、视频"),
            ("第4节 综合案例", "导入机器人、工具和修改TCP操作", "视频"),
        ],
    )

    add_heading(doc, "第四章 工业机器人系统工作轨迹生成", 2)
    add_resource_table(
        doc,
        [
            ("第1节 轨迹获取方式", "工业机器人系统工作轨迹", "PDF"),
            ("第1节 轨迹获取方式", "工业机器人系统工作轨迹优化", "视频"),
            ("第1节 轨迹获取方式", "工业机器人系统工作轨迹生成 - 轨迹优化", "PDF、视频"),
            ("第1节 轨迹获取方式", "工业机器人系统工作轨迹生成 - 曲线特征", "PDF、视频"),
            ("第1节 轨迹获取方式", "工业机器人系统工作轨迹生成 - 一个面的外环", "PDF、视频"),
            ("第2节 轨迹生成", "PQArt软件中打孔轨迹生成操作", "动画"),
            ("第2节 轨迹生成", "PQArt软件中点云打孔轨迹生成操作", "动画"),
            ("第2节 轨迹生成", "PQArt软件中普通打孔轨迹生成", "视频"),
            ("第2节 轨迹生成", "PQArt软件中单条边轨迹生成操作", "动画"),
            ("第2节 轨迹生成", "PQArt软件中曲线特征轨迹生成", "视频"),
            ("第2节 轨迹生成", "PQArt软件中生成轨迹选项_Z向与侧面平行", "视频"),
            ("第2节 轨迹生成", "PQArt软件中生成轨迹选项_反向搜索", "视频"),
            ("第2节 轨迹生成", "PQArt软件中沿着一个面的一条边_终止条件", "视频"),
            ("第2节 轨迹生成", "PQArt软件中沿着一个面的一条边的轨迹生成", "视频"),
            ("第2节 轨迹生成", "PQArt软件中一个面的外环的轨迹生成", "视频"),
            ("第2节 轨迹生成", "PQArt软件中一个面的一个环的轨迹生成", "视频"),
        ],
    )

    add_heading(doc, "第五章 工业机器人离线编程仿真及联机调试", 2)
    add_resource_table(
        doc,
        [
            ("第1节 工业机器人离线编程仿真及联机调试", "离线编程软件联机调试", "PDF"),
            ("第1节 工业机器人离线编程仿真及联机调试", "工业机器人离线编程仿真软件联机调试", "视频"),
            ("第1节 工业机器人离线编程仿真及联机调试", "离线编程仿真操作", "视频"),
            ("第2节 案例：激光切割应用", "工业机器人离线编程仿真案例：激光切割 - 工作站构建", "PDF、视频"),
            ("第2节 案例：激光切割应用", "工业机器人离线编程仿真案例：激光切割 - 生成轨迹", "PDF、视频"),
            ("第2节 案例：激光切割应用", "工业机器人离线编程仿真案例：激光切割 - 轨迹优化", "PDF、视频"),
            ("第2节 案例：激光切割应用", "工业机器人离线编程仿真案例：激光切割 - 仿真和后置", "PDF、视频"),
            ("第2节 案例：激光切割应用", "现场激光切割机器人切割弯管", "视频"),
            ("第2节 案例：激光切割应用", "现场激光切割机器人切割直管", "视频"),
            ("第2节 案例：激光切割应用", "激光切割机器人切割离线编程1", "视频"),
            ("第2节 案例：激光切割应用", "激光切割机器人切割离线编程2", "视频"),
            ("第3节 案例：去毛刺应用", "工业机器人离线编程仿真案例：去毛刺 - 环境搭建", "PDF、视频"),
            ("第3节 案例：去毛刺应用", "工业机器人离线编程仿真案例：去毛刺 - 轨迹生成", "PDF、视频"),
            ("第3节 案例：去毛刺应用", "工业机器人离线编程仿真案例：去毛刺 - 轨迹优化", "PDF、视频"),
            ("第3节 案例：去毛刺应用", "工业机器人离线编程仿真案例：去毛刺 - 后置和仿真", "PDF、视频"),
            ("第3节 案例：去毛刺应用", "去毛刺应用案例", "视频"),
            ("第4节 案例：打孔应用", "工业机器人离线编程仿真案例：打孔 - 环境搭建", "PDF、视频"),
            ("第4节 案例：打孔应用", "工业机器人离线编程仿真案例：打孔 - 轨迹生成", "PDF、视频"),
            ("第4节 案例：打孔应用", "工业机器人离线编程仿真案例：打孔 - 轨迹优化", "PDF、视频"),
            ("第4节 案例：打孔应用", "工业机器人离线编程仿真案例：打孔 - 仿真和后置", "PDF、视频"),
            ("第5节 案例：写字应用", "工业机器人离线编程仿真案例：写字 - 环境搭建", "PDF、视频"),
            ("第5节 案例：写字应用", "工业机器人离线编程仿真案例：写字 - 轨迹生成", "PDF、视频"),
            ("第5节 案例：写字应用", "工业机器人离线编程仿真案例：写字 - 轨迹优化", "PDF、视频"),
            ("第5节 案例：写字应用", "工业机器人离线编程仿真案例：写字 - 仿真和后置", "PDF、视频"),
            ("第5节 案例：写字应用", "PQArt离线编程焊接写字应用", "视频"),
            ("第5节 案例：写字应用", "PQArt离线编程雕刻写字应用", "视频"),
            ("第6节 案例：画画", "工业机器人离线编程仿真案例：画画 - 环境搭建", "PDF、视频"),
            ("第6节 案例：画画", "工业机器人离线编程仿真案例：画画 - 轨迹生成", "PDF、视频"),
            ("第6节 案例：画画", "工业机器人离线编程仿真案例：画画 - 轨迹优化", "PDF、视频"),
            ("第6节 案例：画画", "工业机器人离线编程仿真案例：画画 - 仿真和后置", "PDF、视频"),
        ],
    )

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(16)
    r = note.add_run(
        "说明：源文件为扫描版海报 PDF（2 页），无可选中文字。Word 文档前两页保留原海报版式，"
        "后续页面为可编辑的课程资源清单，软件名称统一为 PQArt。"
    )
    set_run_font(r, size=10, color=RGBColor(0x66, 0x66, 0x66))

    doc.save(OUT_DOCX)
    print(f"Wrote {OUT_DOCX} ({OUT_DOCX.stat().st_size} bytes)")


if __name__ == "__main__":
    build()
